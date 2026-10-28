import json
import os
from typing import Any, Dict, List, Union

class InputProcessor:

    def __init__(self, gemini_api_key: str = None):
        """Khởi tạo InputProcessor sử dụng google.genai SDK mới"""
        if gemini_api_key:
            from google import genai
            self.client = genai.Client(api_key=gemini_api_key)
        else:
            self.client = None

    def _clean_text(self, val: Any) -> Any:
        if isinstance(val, str):
            cleaned = (
                val.replace("\r\n", " ")
                .replace("\n", " ")
                .replace("\r", " ")
                .replace("\t", " ")
            )
            return " ".join(cleaned.split())
        return val

    def _clean_structure(self, data: Any) -> Any:
        if isinstance(data, dict):
            return {
                self._clean_text(k): self._clean_structure(v)
                for k, v in data.items()
            }
        elif isinstance(data, list):
            return [self._clean_structure(item) for item in data]
        else:
            return self._clean_text(data)

    def parse_excel(self, file_path: str) -> Dict[str, Any]:
        try:
            import pandas as pd
            excel_file = pd.ExcelFile(file_path)
            sheets_data = {}

            for sheet_name in excel_file.sheet_names:
                df_raw = pd.read_excel(file_path, sheet_name=sheet_name, header=None)
                df_raw_filled = df_raw.ffill(axis=0).fillna("")

                raw_matrix = df_raw_filled.values.tolist()
                csv_text = df_raw_filled.to_csv(index=False, header=False)

                df_standard = pd.read_excel(file_path, sheet_name=sheet_name)
                df_standard = df_standard.ffill(axis=0).fillna("")

                clean_cols = []
                for i, col in enumerate(df_standard.columns):
                    col_str = self._clean_text(str(col))
                    if not col_str or "Unnamed:" in col_str:
                        col_str = f"Col_{i+1}"
                    clean_cols.append(col_str)

                df_standard.columns = clean_cols
                records_dict = df_standard.to_dict(orient="records")

                clean_records = self._clean_structure(records_dict)
                clean_raw_matrix = self._clean_structure(raw_matrix)

                sheets_data[sheet_name] = {
                    "records": clean_records,
                    "raw_matrix": clean_raw_matrix,
                    "csv_format": csv_text,
                }

            return {
                "file_name": os.path.basename(file_path),
                "file_type": "excel",
                "content": sheets_data,
            }
        except Exception as e:
            return {
                "file_name": os.path.basename(file_path),
                "file_type": "excel",
                "error": str(e),
            }

    def parse_pdf(self, file_path: str) -> Dict[str, Any]:
        try:
            import pdfplumber
            pages_data = []
            with pdfplumber.open(file_path) as pdf:
                for idx, page in enumerate(pdf.pages, 1):
                    page_text = page.extract_text() or ""
                    page_tables = page.extract_tables() or []

                    pages_data.append(
                        {
                            "page": idx,
                            "text": self._clean_text(page_text),
                            "tables": self._clean_structure(page_tables),
                        }
                    )

            return {
                "file_name": os.path.basename(file_path),
                "file_type": "pdf",
                "content": pages_data,
            }
        except Exception as e:
            return {
                "file_name": os.path.basename(file_path),
                "file_type": "pdf",
                "error": str(e),
            }

    def parse_docx(self, file_path: str) -> Dict[str, Any]:
        try:
            from docx import Document
            doc = Document(file_path)
            paragraphs = [
                self._clean_text(p.text)
                for p in doc.paragraphs
                if p.text.strip()
            ]

            tables_data = []
            for table in doc.tables:
                table_rows = []
                for row in table.rows:
                    table_rows.append(
                        [self._clean_text(cell.text) for cell in row.cells]
                    )
                tables_data.append(table_rows)

            return {
                "file_name": os.path.basename(file_path),
                "file_type": "docx",
                "content": {"paragraphs": paragraphs, "tables": tables_data},
            }
        except Exception as e:
            return {
                "file_name": os.path.basename(file_path),
                "file_type": "docx",
                "error": str(e),
            }

    def parse_image(self, file_path: str) -> Dict[str, Any]:
        if not self.client:
            return {
                "file_name": os.path.basename(file_path),
                "file_type": "image",
                "error": "Chưa cấu hình Gemini API Key",
            }

        try:
            uploaded_file = self.client.files.upload(file=file_path)
            prompt = (
                "Hãy trích xuất chính xác tất cả văn bản, ký hiệu kỹ thuật, "
                "bảng dữ liệu và ghi chú trong hình ảnh/bản vẽ này."
            )
            response = self.client.models.generate_content(
                model="gemini-1.5-flash",
                contents=[uploaded_file, prompt]
            )
            return {
                "file_name": os.path.basename(file_path),
                "file_type": "image",
                "content": self._clean_text(response.text),
            }
        except Exception as e:
            return {
                "file_name": os.path.basename(file_path),
                "file_type": "image",
                "error": str(e),
            }

    def process_project_files(
        self, file_paths: List[str], user_prompt: str = ""
    ) -> Dict[str, Any]:
        parsed_files = []
        for path in file_paths:
            ext = os.path.splitext(path)[1].lower()
            if ext in [".xlsx", ".xls", ".csv"]:
                parsed_files.append(self.parse_excel(path))
            elif ext == ".pdf":
                parsed_files.append(self.parse_pdf(path))
            elif ext in [".docx", ".doc"]:
                parsed_files.append(self.parse_docx(path))
            elif ext in [".png", ".jpg", ".jpeg"]:
                parsed_files.append(self.parse_image(path))

        return {
            "session_prompt": user_prompt,
            "project_files": parsed_files,
        }
